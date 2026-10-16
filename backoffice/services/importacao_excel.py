import io
import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any, BinaryIO

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from backoffice.models import (
    Animal, Departamento, EstoqueMovimento, GrupoProduto, ImportacaoPlanilha, Marca, Produto, ProdutoImagem,
    Subdepartamento,
)
from backoffice.services.texto import normalizar, slugificar

from .canais import definir_preco_site, garantir_canais
from .estoque import definir_saldo, definir_seguranca

logger = logging.getLogger(__name__)

LIMITE_LINHAS = 5000
OBRIGATORIAS = ("sku", "nome", "animal", "preco")
COLUNAS = (
    "sku", "nome", "descricao", "marca", "animal", "departamento", "subdepartamento", "grupo", "variacao",
    "gtin", "ncm", "peso_g", "altura_cm", "largura_cm", "comprimento_cm", "custo",
    "preco", "preco_promocional", "promo_inicio", "promo_fim", "estoque", "estoque_seguranca",
    "ativo", "publicar_site", "tipo_arte", "cor_arte", "imagem_url", "imagens",
)
FORMATOS_DE_DATA = ("%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M", "%Y-%m-%d")
SEPARADOR_IMAGENS = "|"
PADRAO_GTIN = re.compile(r"^(\d{8}|\d{12,14})$")
PADRAO_NCM = re.compile(r"^\d{8}$")
APELIDOS = {
    "promocao": "preco_promocional", "publicar_no_site": "publicar_site",
    "ean": "gtin", "gtin_ean": "gtin", "quantidade": "estoque",
}
VERDADEIRO = {"sim", "s", "1", "true", "verdadeiro", "x", "yes"}
FALSO = {"nao", "n", "0", "false", "falso", "no"}


class PlanilhaInvalida(Exception):
    pass


@dataclass(frozen=True)
class ErroLinha:
    linha: int
    campo: str
    mensagem: str


@dataclass
class ResultadoImportacao:
    total: int = 0
    criados: int = 0
    atualizados: int = 0
    erros: list[ErroLinha] = field(default_factory=list)


class _Taxonomia:
    def __init__(self):
        self.animais = self._indexar(Animal.objects.all())
        self.departamentos = self._indexar(Departamento.objects.all())
        self.subdepartamentos = self._indexar(Subdepartamento.objects.select_related("departamento"))

    @staticmethod
    def _indexar(objetos) -> dict[str, Any]:
        indice = {}
        for objeto in objetos:
            indice[normalizar(objeto.slug)] = objeto
            indice[normalizar(objeto.nome)] = objeto
        return indice

    def buscar(self, indice: dict[str, Any], valor: str, campo: str):
        objeto = indice.get(normalizar(valor))
        if objeto is None:
            raise ValueError(f"{campo} '{valor}' não cadastrado.")
        return objeto


def importar_planilha(
    arquivo: BinaryIO, *, nome_arquivo: str, usuario: str, apenas_validar: bool = False,
) -> ResultadoImportacao:
    linhas = _ler_linhas(arquivo)
    resultado = ResultadoImportacao(total=len(linhas))
    taxonomia = _Taxonomia()

    with transaction.atomic():
        for numero, valores in linhas:
            try:
                with transaction.atomic():
                    criado = _aplicar_linha(valores, taxonomia, referencia=nome_arquivo)
            except ValueError as erro:
                resultado.erros.append(ErroLinha(numero, _campo_do_erro(str(erro)), str(erro)))
            except Exception as erro:  # noqa: BLE001 — uma linha ruim não derruba a planilha inteira
                logger.exception("importacao_linha_falhou linha=%d", numero)
                resultado.erros.append(ErroLinha(numero, "", f"Erro inesperado: {erro}"))
            else:
                if criado:
                    resultado.criados += 1
                else:
                    resultado.atualizados += 1
        if apenas_validar:
            transaction.set_rollback(True)

    if not apenas_validar:
        ImportacaoPlanilha.objects.create(
            arquivo=nome_arquivo[:200], usuario=usuario, total=resultado.total, criados=resultado.criados,
            atualizados=resultado.atualizados, erros=[erro.__dict__ for erro in resultado.erros],
        )
    logger.info(
        "planilha_importada arquivo=%s validar=%s total=%d criados=%d atualizados=%d erros=%d",
        nome_arquivo, apenas_validar, resultado.total, resultado.criados, resultado.atualizados, len(resultado.erros),
    )
    return resultado


def gerar_modelo_planilha() -> bytes:
    livro = Workbook()
    produtos = livro.active
    produtos.title = "Produtos"
    produtos.append(list(COLUNAS))
    produtos.append([
        "RACAO-015", "Ração Premium Adultos 15kg", "Alimento completo para cães adultos.", "NutriPet", "cães", "Rações",
        "Ração Seca", "Ração Premium Adultos", "15 kg", "7891234567890", "23091000", 15200, 62, 40, 20, "120,00",
        "189,90", "159,90", "10/10/2026", "20/10/2026 23:59", 40, 2, "sim", "sim", "saco", "#D96C4A", "",
        "https://exemplo.com/foto-1.jpg|https://exemplo.com/foto-2.jpg",
    ])
    instrucoes = livro.create_sheet("Instruções")
    for linha in (
        ["Coluna", "Regra"],
        ["sku, nome, animal, preco", "Obrigatórias. O SKU identifica o produto: se já existir, ele é atualizado."],
        ["animal, departamento, subdepartamento", "Use o nome ou o código (slug) já cadastrado. Marca nova é criada sozinha."],
        ["preco, preco_promocional, custo", "Aceita 12,90 ou 12.90. O promocional não pode passar do preço. Valem para o site."],
        ["estoque", "Saldo desejado. O sistema registra a diferença como movimento de importação."],
        ["ativo, publicar_site", "sim ou não. Em branco: produto novo fica ativo e publicado."],
        ["grupo, variacao", "Produtos com o mesmo grupo viram variações um do outro (ex.: 3 kg e 15 kg). A variação é o rótulo."],
        ["gtin, ncm, peso_g, altura_cm, largura_cm, comprimento_cm", "Exigidos para publicar em marketplaces. GTIN com 8 ou 12 a 14 dígitos; NCM com 8 dígitos; peso em gramas; medidas em cm."],
        ["promo_inicio, promo_fim", "Datas da promoção (dd/mm/aaaa ou dd/mm/aaaa hh:mm). Sem datas, a promoção vale até você tirá-la."],
        ["estoque_seguranca", "Unidades que nunca são vendidas online."],
        ["imagens", "Endereços das fotos separados por | (a primeira é a principal)."],
        ["tipo_arte", "saco, lata, frasco, bola, osso, caixa, cama, coleira ou pote."],
        ["Células vazias", "Em produto existente, célula vazia mantém o valor atual."],
    ):
        instrucoes.append(linha)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _ler_linhas(arquivo: BinaryIO) -> list[tuple[int, dict[str, Any]]]:
    try:
        livro = load_workbook(arquivo, read_only=True, data_only=True)
    except Exception as erro:  # noqa: BLE001 — qualquer falha de leitura vira erro de planilha
        raise PlanilhaInvalida("Não foi possível ler o arquivo. Envie um .xlsx válido.") from erro
    folha = livro.worksheets[0]
    linhas = folha.iter_rows(values_only=True)
    cabecalho = next(linhas, None)
    if cabecalho is None:
        raise PlanilhaInvalida("A planilha está vazia.")
    colunas = [_normalizar_coluna(valor) for valor in cabecalho]
    faltando = [c for c in OBRIGATORIAS if c not in colunas]
    if faltando:
        raise PlanilhaInvalida(f"Faltam as colunas obrigatórias: {', '.join(faltando)}.")

    registros = []
    for numero, valores in enumerate(linhas, start=2):
        registro = {coluna: valor for coluna, valor in zip(colunas, valores) if coluna in COLUNAS}
        if all(v is None or str(v).strip() == "" for v in registro.values()):
            continue
        registros.append((numero, registro))
    if len(registros) > LIMITE_LINHAS:
        raise PlanilhaInvalida(f"A planilha passa do limite de {LIMITE_LINHAS} linhas.")
    return registros


def _normalizar_coluna(valor: Any) -> str:
    texto = normalizar(str(valor or "")).replace(" ", "_")
    return APELIDOS.get(texto, texto)


def _campo_do_erro(mensagem: str) -> str:
    return max((campo for campo in COLUNAS if mensagem.startswith(campo)), key=len, default="")


def _texto(valores: dict[str, Any], campo: str) -> str:
    valor = valores.get(campo)
    return "" if valor is None else str(valor).strip()


def _decimal(valores: dict[str, Any], campo: str) -> Decimal | None:
    bruto = valores.get(campo)
    if bruto is None or str(bruto).strip() == "":
        return None
    texto = str(bruto).strip().replace("R$", "").replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        numero = Decimal(texto).quantize(Decimal("0.01"))
    except InvalidOperation as erro:
        raise ValueError(f"{campo}: valor '{bruto}' não é um número.") from erro
    if numero < 0:
        raise ValueError(f"{campo}: não pode ser negativo.")
    return numero


def _inteiro(valores: dict[str, Any], campo: str) -> int | None:
    bruto = valores.get(campo)
    if bruto is None or str(bruto).strip() == "":
        return None
    try:
        numero = int(Decimal(str(bruto).replace(",", ".")))
    except (InvalidOperation, ValueError) as erro:
        raise ValueError(f"{campo}: valor '{bruto}' não é um número inteiro.") from erro
    if numero < 0:
        raise ValueError(f"{campo}: não pode ser negativo.")
    return numero


def _data(valores: dict[str, Any], campo: str) -> datetime | None:
    bruto = valores.get(campo)
    if bruto is None or str(bruto).strip() == "":
        return None
    if isinstance(bruto, datetime):
        data = bruto
    else:
        for formato in FORMATOS_DE_DATA:
            try:
                data = datetime.strptime(str(bruto).strip(), formato)
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"{campo}: data '{bruto}' inválida. Use dd/mm/aaaa ou dd/mm/aaaa hh:mm.")
    return timezone.make_aware(data) if timezone.is_naive(data) else data


def _booleano(valores: dict[str, Any], campo: str) -> bool | None:
    bruto = valores.get(campo)
    if bruto is None or str(bruto).strip() == "":
        return None
    texto = normalizar(str(bruto))
    if texto in VERDADEIRO:
        return True
    if texto in FALSO:
        return False
    raise ValueError(f"{campo}: use sim ou não (recebido '{bruto}').")


def _aplicar_linha(valores: dict[str, Any], taxonomia: _Taxonomia, referencia: str) -> bool:
    for campo in OBRIGATORIAS:
        if not _texto(valores, campo):
            raise ValueError(f"{campo}: campo obrigatório vazio.")

    sku = _texto(valores, "sku")
    produto = Produto.objects.filter(sku=sku).first()
    criado = produto is None
    if criado:
        produto = Produto(sku=sku, slug=_slug_unico(_texto(valores, "nome")))

    produto.nome = _texto(valores, "nome")
    produto.animal = taxonomia.buscar(taxonomia.animais, _texto(valores, "animal"), "animal")
    _aplicar_departamentos(produto, valores, taxonomia)
    _aplicar_opcionais(produto, valores)

    preco, promo = _decimal(valores, "preco"), _decimal(valores, "preco_promocional")
    if preco is None:
        raise ValueError("preco: campo obrigatório vazio.")
    if promo is not None and promo > preco:
        raise ValueError("preco_promocional: não pode ser maior que o preço.")
    inicio, fim = _data(valores, "promo_inicio"), _data(valores, "promo_fim")
    if inicio and fim and fim <= inicio:
        raise ValueError("promo_fim: precisa ser depois de promo_inicio.")
    produto.preco = preco
    _aplicar_grupo(produto, valores)
    produto.save()

    publicar = _booleano(valores, "publicar_site")
    existente = None if criado else produto.canais.filter(canal__codigo="site").first()
    publicado = publicar if publicar is not None else (True if existente is None else existente.publicado)
    garantir_canais(produto, preco, promo, publicar_site=publicado)
    linha = definir_preco_site(produto, preco, promo, publicado)
    if inicio or fim or "promo_inicio" in valores or "promo_fim" in valores:
        linha.promo_inicio, linha.promo_fim = inicio, fim
        linha.save(update_fields=["promo_inicio", "promo_fim"])
    _aplicar_imagens(produto, valores)

    estoque = _inteiro(valores, "estoque")
    if estoque is not None:
        definir_saldo(produto, estoque, EstoqueMovimento.Motivo.IMPORTACAO, referencia[:80])
    seguranca = _inteiro(valores, "estoque_seguranca")
    if seguranca is not None:
        definir_seguranca(produto, seguranca)
    return criado


def _aplicar_departamentos(produto: Produto, valores: dict[str, Any], taxonomia: _Taxonomia) -> None:
    nome_sub = _texto(valores, "subdepartamento")
    nome_dep = _texto(valores, "departamento")
    if nome_sub:
        sub = taxonomia.buscar(taxonomia.subdepartamentos, nome_sub, "subdepartamento")
        if nome_dep and taxonomia.buscar(taxonomia.departamentos, nome_dep, "departamento").pk != sub.departamento_id:
            raise ValueError(f"subdepartamento: '{nome_sub}' não pertence ao departamento '{nome_dep}'.")
        produto.subdepartamento, produto.departamento = sub, sub.departamento
    elif nome_dep:
        produto.departamento = taxonomia.buscar(taxonomia.departamentos, nome_dep, "departamento")
        produto.subdepartamento = None


def _aplicar_grupo(produto: Produto, valores: dict[str, Any]) -> None:
    nome, variacao = _texto(valores, "grupo"), _texto(valores, "variacao")
    if nome:
        produto.grupo, _ = GrupoProduto.objects.get_or_create(slug=slugificar(nome), defaults={"nome": nome})
    if variacao:
        if produto.grupo_id is None:
            raise ValueError("variacao: informe também o grupo.")
        produto.variacao = variacao


def _aplicar_imagens(produto: Produto, valores: dict[str, Any]) -> None:
    urls = [u.strip() for u in _texto(valores, "imagens").split(SEPARADOR_IMAGENS) if u.strip()]
    if not urls:
        return
    ProdutoImagem.objects.filter(produto=produto).delete()
    ProdutoImagem.objects.bulk_create(ProdutoImagem(produto=produto, url=url[:500], posicao=i) for i, url in enumerate(urls))


def _aplicar_opcionais(produto: Produto, valores: dict[str, Any]) -> None:
    for campo in ("descricao", "imagem_url", "cor_arte"):
        if _texto(valores, campo):
            setattr(produto, campo, _texto(valores, campo))
    for campo, padrao, dica in (("gtin", PADRAO_GTIN, "8 ou 12 a 14 dígitos"), ("ncm", PADRAO_NCM, "8 dígitos")):
        texto = _texto(valores, campo)
        if texto:
            if not padrao.match(texto):
                raise ValueError(f"{campo}: '{texto}' inválido, use {dica}, só números.")
            setattr(produto, campo, texto)
    peso = _inteiro(valores, "peso_g")
    if peso is not None:
        if peso == 0:
            raise ValueError("peso_g: precisa ser maior que zero.")
        produto.peso_g = peso
    for campo in ("altura_cm", "largura_cm", "comprimento_cm"):
        medida = _decimal(valores, campo)
        if medida is not None:
            if medida == 0:
                raise ValueError(f"{campo}: precisa ser maior que zero.")
            setattr(produto, campo, medida)
    tipo = normalizar(_texto(valores, "tipo_arte"))
    if tipo:
        if tipo not in Produto.TipoArte.values:
            raise ValueError(f"tipo_arte: '{tipo}' inválido. Use {', '.join(Produto.TipoArte.values)}.")
        produto.tipo_arte = tipo
    nome_marca = _texto(valores, "marca")
    if nome_marca:
        produto.marca, _ = Marca.objects.get_or_create(nome=nome_marca, defaults={"slug": slugificar(nome_marca)})
    custo = _decimal(valores, "custo")
    if custo is not None:
        produto.custo = custo
    ativo = _booleano(valores, "ativo")
    if ativo is not None:
        produto.ativo = ativo
    elif produto.pk is None:
        produto.ativo = True


def _slug_unico(nome: str) -> str:
    base = slugificar(nome)[:150] or "produto"
    slug, contador = base, 2
    while Produto.objects.filter(slug=slug).exists():
        slug, contador = f"{base}-{contador}", contador + 1
    return slug
