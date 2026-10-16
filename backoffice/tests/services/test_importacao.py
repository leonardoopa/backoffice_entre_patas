from decimal import Decimal

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from datetime import timedelta

from backoffice.models import GrupoProduto, ImportacaoPlanilha, Marca, Produto, ProdutoCanal, ProdutoImagem
from backoffice.services.estoque import saldo_atual
from backoffice.services.importacao_excel import COLUNAS, PlanilhaInvalida, gerar_modelo_planilha, importar_planilha

from ..support.fabricas import criar_catalogo, criar_taxonomia, planilha


def importar(linhas, colunas=None, apenas_validar=False):
    return importar_planilha(planilha(linhas, colunas), nome_arquivo="t.xlsx", usuario="ana", apenas_validar=apenas_validar)


class ImportacaoTests(TestCase):
    def setUp(self):
        criar_taxonomia()

    def test_cria_produto_com_preco_estoque_marca_e_canais(self):
        r = importar([["N-1", "Novo Produto", "cães", "1.234,56", "999,90", 12, "Rações", "Ração Seca", "Marca Nova"]])
        self.assertEqual((r.total, r.criados, r.atualizados, r.erros), (1, 1, 0, []))
        produto = Produto.objects.get(sku="N-1")
        self.assertEqual((produto.slug, produto.preco, produto.animal.slug), ("novo-produto", Decimal("1234.56"), "caes"))
        self.assertEqual((produto.departamento.slug, produto.subdepartamento.slug, produto.marca.nome), ("racoes", "racao-seca", "Marca Nova"))
        self.assertEqual(saldo_atual(produto), 12)
        site = ProdutoCanal.objects.get(produto=produto, canal__codigo="site")
        self.assertEqual((site.preco, site.preco_promocional, site.publicado), (Decimal("1234.56"), Decimal("999.90"), True))
        marketplaces = ProdutoCanal.objects.filter(produto=produto).exclude(canal__codigo="site")
        self.assertEqual((marketplaces.count(), marketplaces.filter(publicado=True).count()), (2, 0))

    def test_atualiza_produto_existente_pelo_sku(self):
        importar([["N-1", "Antigo", "caes", "10", "", 5, "", "", ""]])
        r = importar([["N-1", "Novo nome", "caes", "20,50", "", 9, "", "", ""]])
        self.assertEqual((r.criados, r.atualizados), (0, 1))
        produto = Produto.objects.get(sku="N-1")
        self.assertEqual((produto.nome, produto.preco, saldo_atual(produto), Produto.objects.count()), ("Novo nome", Decimal("20.50"), 9, 1))
        self.assertEqual(ProdutoCanal.objects.get(produto=produto, canal__codigo="site").preco, Decimal("20.50"))

    def test_celula_vazia_mantem_o_valor_atual(self):
        importar([["N-1", "Produto", "caes", "10", "", 5, "Rações", "Ração Seca", "Marca A"]])
        importar([["N-1", "Produto", "caes", "10", "", "", "", "", ""]])
        produto = Produto.objects.get(sku="N-1")
        self.assertEqual((produto.departamento.slug, produto.marca.nome, saldo_atual(produto)), ("racoes", "Marca A", 5))

    def test_slug_repetido_ganha_sufixo(self):
        r = importar([["A", "Mesmo Nome", "caes", "1", "", "", "", "", ""], ["B", "Mesmo Nome", "caes", "1", "", "", "", "", ""]])
        self.assertEqual(r.criados, 2)
        self.assertEqual(set(Produto.objects.values_list("slug", flat=True)), {"mesmo-nome", "mesmo-nome-2"})

    def test_erros_por_linha_nao_derrubam_as_linhas_boas(self):
        r = importar([
            ["OK-1", "Bom", "caes", "10", "", "", "", "", ""],
            ["ERR-1", "Animal ruim", "dragoes", "10", "", "", "", "", ""],
            ["ERR-2", "Preço ruim", "caes", "abc", "", "", "", "", ""],
            ["ERR-3", "Promo maior", "caes", "10", "11", "", "", "", ""],
            ["ERR-4", "Sub fora do dep", "caes", "10", "", "", "Brinquedos", "Ração Seca", ""],
            ["", "Sem sku", "caes", "10", "", "", "", "", ""],
        ])
        self.assertEqual((r.criados, len(r.erros)), (1, 5))
        self.assertEqual([e.linha for e in r.erros], [3, 4, 5, 6, 7])
        self.assertEqual({e.campo for e in r.erros}, {"animal", "preco", "preco_promocional", "subdepartamento", "sku"})
        self.assertEqual(list(Produto.objects.values_list("sku", flat=True)), ["OK-1"])

    def test_apenas_validar_nao_grava_nada(self):
        r = importar([["N-1", "Novo", "caes", "10", "", 5, "", "", ""]], apenas_validar=True)
        self.assertEqual((r.criados, Produto.objects.count(), ImportacaoPlanilha.objects.count()), (1, 0, 0))

    def test_importacao_fica_registrada(self):
        importar([["N-1", "Novo", "caes", "10", "", "", "", "", ""], ["X", "Ruim", "dragoes", "1", "", "", "", "", ""]])
        registro = ImportacaoPlanilha.objects.get()
        self.assertEqual((registro.usuario, registro.total, registro.criados, len(registro.erros)), ("ana", 2, 1, 1))

    def test_cabecalho_aceita_apelidos_e_acentos(self):
        r = importar([["N-1", "Novo", "caes", "10", 7]], ["SKU", "Nome", "Animal", "Preço", "Quantidade"])
        self.assertEqual(r.criados, 1)
        self.assertEqual(saldo_atual(Produto.objects.get(sku="N-1")), 7)

    def test_linhas_em_branco_sao_ignoradas(self):
        r = importar([["N-1", "Novo", "caes", "10"], [None, None, None, None], ["N-2", "Outro", "caes", "10"]], ["sku", "nome", "animal", "preco"])
        self.assertEqual((r.total, r.criados), (2, 2))

    def test_booleanos_em_portugues(self):
        importar([["N-1", "Novo", "caes", "10", "não", "não"]], ["sku", "nome", "animal", "preco", "ativo", "publicar_site"])
        produto = Produto.objects.get(sku="N-1")
        self.assertFalse(produto.ativo)
        self.assertFalse(ProdutoCanal.objects.get(produto=produto, canal__codigo="site").publicado)

    def test_colunas_obrigatorias_ausentes(self):
        with self.assertRaisesMessage(PlanilhaInvalida, "preco"):
            importar([["N-1", "Novo", "caes"]], ["sku", "nome", "animal"])

    def test_arquivo_que_nao_e_xlsx(self):
        import io

        with self.assertRaises(PlanilhaInvalida):
            importar_planilha(io.BytesIO(b"isto nao e excel"), nome_arquivo="x.xlsx", usuario="ana")

    def test_modelo_da_planilha_e_valido_e_importavel(self):
        import io

        conteudo = gerar_modelo_planilha()
        livro = load_workbook(io.BytesIO(conteudo))
        self.assertEqual([c.value for c in livro["Produtos"][1]], list(COLUNAS))
        self.assertIn("Instruções", livro.sheetnames)
        resultado = importar_planilha(io.BytesIO(conteudo), nome_arquivo="modelo.xlsx", usuario="ana")
        self.assertEqual((resultado.criados, resultado.erros), (1, []))

    def test_marca_nao_duplica(self):
        importar([["A", "P1", "caes", "1", "", "", "", "", "Mesma"], ["B", "P2", "caes", "1", "", "", "", "", "mesma"]])
        self.assertEqual(Marca.objects.count(), 1)


class ImportacaoColunasNovasTests(TestCase):
    COLUNAS = ["sku", "nome", "animal", "preco"]

    def setUp(self):
        criar_taxonomia()

    def importar(self, extras, linha_extra):
        return importar([["N-1", "Novo", "caes", "100"] + linha_extra], self.COLUNAS + extras)

    def test_dados_de_marketplace(self):
        r = self.importar(
            ["gtin", "ncm", "peso_g", "altura_cm", "largura_cm", "comprimento_cm"],
            ["7891234567895", "23091000", 15200, "62,5", 40, "20"],
        )
        self.assertEqual(r.erros, [])
        p = Produto.objects.get()
        self.assertEqual((p.gtin, p.ncm, p.peso_g, p.altura_cm, p.largura_cm, p.comprimento_cm), ("7891234567895", "23091000", 15200, Decimal("62.5"), Decimal("40.0"), Decimal("20.0")))

    def test_valores_invalidos_geram_erro_claro(self):
        for extras, valores, campo in (
            (["gtin"], ["123"], "gtin"), (["ncm"], ["12"], "ncm"), (["peso_g"], [0], "peso_g"), (["altura_cm"], [0], "altura_cm"),
            (["promo_inicio"], ["ontem"], "promo_inicio"),
        ):
            r = importar([["X", "Produto", "caes", "100"] + valores], self.COLUNAS + extras)
            self.assertEqual([e.campo for e in r.erros], [campo], extras)
        self.assertEqual(Produto.objects.count(), 0)

    def test_grupo_e_variacao_ligam_produtos(self):
        importar([
            ["A-15", "Ração 15kg", "caes", "100", "Ração Premium", "15 kg"],
            ["A-3", "Ração 3kg", "caes", "40", "Ração Premium", "3 kg"],
        ], self.COLUNAS + ["grupo", "variacao"])
        grupo = GrupoProduto.objects.get()
        self.assertEqual(set(grupo.produtos.values_list("variacao", flat=True)), {"15 kg", "3 kg"})

    def test_variacao_sem_grupo_e_erro(self):
        r = self.importar(["variacao"], ["15 kg"])
        self.assertEqual([e.campo for e in r.erros], ["variacao"])

    def test_variacao_repetida_no_mesmo_grupo_vira_erro_da_linha(self):
        r = importar([
            ["A", "P1", "caes", "10", "G", "1 kg"], ["B", "P2", "caes", "10", "G", "1 kg"],
        ], self.COLUNAS + ["grupo", "variacao"])
        self.assertEqual((r.criados, len(r.erros)), (1, 1))

    def test_datas_da_promocao_em_texto_e_a_vigencia(self):
        r = importar([["N-1", "Novo", "caes", "100", "80", "10/10/2026", "20/10/2026 23:59"]], self.COLUNAS + ["preco_promocional", "promo_inicio", "promo_fim"])
        self.assertEqual(r.erros, [])
        linha = ProdutoCanal.objects.get(canal__codigo="site")
        inicio, fim = timezone.localtime(linha.promo_inicio), timezone.localtime(linha.promo_fim)
        self.assertEqual((inicio.day, inicio.month, fim.day, fim.hour, fim.minute), (10, 10, 20, 23, 59))

    def test_fim_antes_do_inicio_e_erro(self):
        r = importar([["N-1", "Novo", "caes", "100", "80", "20/10/2026", "10/10/2026"]], self.COLUNAS + ["preco_promocional", "promo_inicio", "promo_fim"])
        self.assertEqual([e.campo for e in r.erros], ["promo_fim"])

    def test_datas_vazias_removem_o_prazo_da_promocao_existente(self):
        importar([["N-1", "Novo", "caes", "100", "80", "10/10/2026", "20/10/2026"]], self.COLUNAS + ["preco_promocional", "promo_inicio", "promo_fim"])
        importar([["N-1", "Novo", "caes", "100", "80", "", ""]], self.COLUNAS + ["preco_promocional", "promo_inicio", "promo_fim"])
        linha = ProdutoCanal.objects.get(canal__codigo="site")
        self.assertEqual((linha.promo_inicio, linha.promo_fim), (None, None))

    def test_imagens_separadas_por_barra_substituem_as_anteriores(self):
        self.importar(["imagens"], ["https://x/1.jpg|https://x/2.jpg"])
        self.assertEqual(list(ProdutoImagem.objects.order_by("posicao").values_list("url", flat=True)), ["https://x/1.jpg", "https://x/2.jpg"])
        self.importar(["imagens"], ["https://x/3.jpg"])
        self.assertEqual(list(ProdutoImagem.objects.values_list("url", flat=True)), ["https://x/3.jpg"])

    def test_estoque_de_seguranca(self):
        self.importar(["estoque", "estoque_seguranca"], [10, 3])
        estoque = Produto.objects.get().estoque
        self.assertEqual((estoque.quantidade, estoque.seguranca, estoque.disponivel), (10, 3, 7))
